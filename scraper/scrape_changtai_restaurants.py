"""
上海长泰广场（浦东新区张江镇）500m 内餐饮商户爬取
数据来源：OpenStreetMap Overpass API（无需 API Key）
输出：Excel 文件
"""

import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime
import math

# 长泰广场（张江镇）坐标 —— 如有偏差可在高德地图查询后替换
CENTER_LAT = 31.2047
CENTER_LNG = 121.6060
RADIUS_M = 500

OVERPASS_URL = "https://overpass-api.de/api/interpreter"

FOOD_AMENITIES = [
    "restaurant", "cafe", "fast_food", "food_court",
    "bar", "pub", "ice_cream", "bakery", "canteen", "bubble_tea"
]

def build_overpass_query(lat, lng, radius):
    amenity_filter = "|".join(FOOD_AMENITIES)
    return f"""
[out:json][timeout:30];
(
  node["amenity"~"{amenity_filter}"](around:{radius},{lat},{lng});
  way["amenity"~"{amenity_filter}"](around:{radius},{lat},{lng});
  node["shop"~"bakery|butcher|deli|food"](around:{radius},{lat},{lng});
);
out center;
""".strip()

def haversine(lat1, lng1, lat2, lng2):
    R = 6371000
    phi1, phi2 = math.radians(lat1), math.radians(lat2)
    dphi = math.radians(lat2 - lat1)
    dlambda = math.radians(lng2 - lng1)
    a = math.sin(dphi/2)**2 + math.cos(phi1)*math.cos(phi2)*math.sin(dlambda/2)**2
    return round(2 * R * math.asin(math.sqrt(a)))

def fetch_pois():
    query = build_overpass_query(CENTER_LAT, CENTER_LNG, RADIUS_M)
    print("正在查询 OpenStreetMap Overpass API...")
    for url in [OVERPASS_URL, "https://overpass.kumi.systems/api/interpreter"]:
        try:
            resp = requests.post(url, data={"data": query}, timeout=40)
            resp.raise_for_status()
            return resp.json().get("elements", [])
        except Exception as e:
            print(f"  节点 {url} 失败: {e}")
    return []

def parse_element(el):
    tags = el.get("tags", {})
    if el["type"] == "node":
        lat, lng = el.get("lat"), el.get("lon")
    else:
        center = el.get("center", {})
        lat, lng = center.get("lat"), center.get("lon")
    if not lat or not lng:
        return None

    name = tags.get("name") or tags.get("name:zh") or tags.get("name:en") or "(无名称)"
    address = " ".join(filter(None, [
        tags.get("addr:street", ""),
        tags.get("addr:housenumber", ""),
        tags.get("addr:full", "")
    ])).strip()

    return {
        "名称":        name,
        "类别":        tags.get("amenity") or tags.get("shop") or "",
        "菜系":        tags.get("cuisine", ""),
        "地址":        address,
        "楼层":        tags.get("level", ""),
        "电话":        tags.get("phone") or tags.get("contact:phone") or "",
        "营业时间":    tags.get("opening_hours", ""),
        "网址":        tags.get("website") or tags.get("contact:website") or "",
        "纬度":        lat,
        "经度":        lng,
        "距中心(m)":   haversine(CENTER_LAT, CENTER_LNG, lat, lng),
        "OSM ID":      f"{el['type']}/{el['id']}",
    }

def save_excel(rows, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "长泰广场餐饮商户"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    even_fill = PatternFill(start_color="EBF3FB", end_color="EBF3FB", fill_type="solid")

    headers = list(rows[0].keys())
    ws.row_dimensions[1].height = 28
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for r, row in enumerate(rows, 2):
        ws.row_dimensions[r].height = 18
        for c, key in enumerate(headers, 1):
            cell = ws.cell(row=r, column=c, value=row.get(key, ""))
            cell.font = Font(size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if r % 2 == 0:
                cell.fill = even_fill

    col_widths = {
        "名称": 28, "类别": 14, "菜系": 14, "地址": 36, "楼层": 8,
        "电话": 16, "营业时间": 24, "网址": 30,
        "纬度": 12, "经度": 12, "距中心(m)": 12, "OSM ID": 20,
    }
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = col_widths.get(h, 14)

    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("说明")
    info = [
        ("数据说明", True),
        ("", False),
        (f"搜索中心：上海长泰广场（张江镇）", False),
        (f"坐标：{CENTER_LAT}, {CENTER_LNG}", False),
        (f"搜索半径：{RADIUS_M} 米", False),
        (f"数据来源：OpenStreetMap Overpass API", False),
        (f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", False),
        (f"商户总数：{len(rows)} 条", False),
        ("", False),
        ("注意事项", True),
        ("1. OSM 数据依赖社区贡献，中国地区覆盖可能不全", False),
        ("2. 如结果偏少，建议改用高德地图 API（免费注册 lbs.amap.com）", False),
        ("3. SAP Vendor Code 匹配需结合内部采购系统另行核对", False),
    ]
    for r, (text, bold) in enumerate(info, 1):
        cell = ws2.cell(row=r, column=1, value=text)
        cell.font = Font(bold=bold, size=11 if bold else 10)
    ws2.column_dimensions["A"].width = 60

    wb.save(filename)

def main():
    print(f"搜索范围：上海长泰广场（张江镇）半径 {RADIUS_M}m")
    print(f"中心坐标：{CENTER_LAT}, {CENTER_LNG}\n")

    elements = fetch_pois()
    print(f"Overpass 返回原始数据：{len(elements)} 条")

    rows, seen = [], set()
    for el in elements:
        parsed = parse_element(el)
        if parsed and parsed["名称"] not in seen:
            seen.add(parsed["名称"])
            rows.append(parsed)

    rows.sort(key=lambda x: x["距中心(m)"])
    print(f"去重后有效商户：{len(rows)} 条\n")

    if not rows:
        print("未获取到数据，可能原因：")
        print("  1. 张江商圈 OSM 录入不全（该区域以办公楼为主，餐饮 POI 录入较少）")
        print("  2. 网络访问 overpass-api.de 受限")
        print("\n建议：改用高德地图 API（免费，lbs.amap.com 注册约 5 分钟）")
        return

    filename = f"长泰广场餐饮商户_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    save_excel(rows, filename)
    print(f"已保存：{filename}（共 {len(rows)} 条）")

if __name__ == "__main__":
    main()
